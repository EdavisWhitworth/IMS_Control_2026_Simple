"""Raw vs processed export/import: raw data must survive a round trip unchanged."""

import numpy as np
import pytest

from ims_control.acquisition.experiment import ExperimentConfig
from ims_control.io.csv_io import export_spectra_csv, import_spectra_csv
from ims_control.io.hdf5_io import export_hdf5, import_hdf5
from ims_control.io.mzml_io import export_mzml, import_mzml
from ims_control.models.data_store import DataStore


def _processed_store():
    """Store whose processed `intensity` differs from `raw_intensity` (flip + scale + offset)."""
    store = DataStore(ExperimentConfig(exp_length_ms=5.0, num_points=50, iterations=2))
    rng = np.random.default_rng(0)
    for i in range(2):
        rec = store.add_iteration(i, rng.normal(size=50) + 3.0)
        rec.intensity = (-rec.raw_intensity - 1.0) / 7.0
    return store


def _assert_raw_and_processed_preserved(loaded, store, exact=True):
    for a, b in zip(loaded.iterations, store.iterations):
        if exact:
            np.testing.assert_array_equal(a.raw_intensity, b.raw_intensity)
        else:  # Excel stores ~15 significant digits
            np.testing.assert_allclose(a.raw_intensity, b.raw_intensity, rtol=1e-14, atol=0)
        np.testing.assert_allclose(a.intensity, b.intensity)
    assert not np.allclose(loaded.iterations[0].raw_intensity, loaded.iterations[0].intensity)


def test_hdf5_default_exports_both_labelled(tmp_path):
    import h5py

    store = _processed_store()
    path = tmp_path / "run.h5"
    export_hdf5(store, path)
    with h5py.File(path, "r") as f:
        assert "raw_intensity" in f and "intensity" in f
        assert f["raw_intensity"].attrs["label"].startswith("raw")
        assert f["intensity"].attrs["label"].startswith("processed")
    _assert_raw_and_processed_preserved(import_hdf5(path), store)


def test_csv_round_trip_preserves_raw_from_either_file(tmp_path):
    store = _processed_store()
    export_spectra_csv(store, tmp_path / "run.csv")
    for name in ("run_processed.csv", "run_raw.csv"):
        loaded = import_spectra_csv(tmp_path / name)
        _assert_raw_and_processed_preserved(loaded, store)


def test_mzml_round_trip_preserves_raw(tmp_path):
    store = _processed_store()
    export_mzml(store, tmp_path / "run.mzml")
    _assert_raw_and_processed_preserved(import_mzml(tmp_path / "run_processed.mzml"), store)


def test_hdf5_selection_raw_only_and_processed_only(tmp_path):
    import h5py

    store = _processed_store()
    export_hdf5(store, tmp_path / "raw.h5", data="raw")
    export_hdf5(store, tmp_path / "proc.h5", data="processed")
    with h5py.File(tmp_path / "raw.h5", "r") as f:
        assert "raw_intensity" in f and "intensity" not in f
    with h5py.File(tmp_path / "proc.h5", "r") as f:
        assert "intensity" in f and "raw_intensity" not in f
    loaded_raw = import_hdf5(tmp_path / "raw.h5")
    np.testing.assert_array_equal(loaded_raw.iterations[0].raw_intensity, store.iterations[0].raw_intensity)


def test_csv_and_mzml_selection_writes_only_selected(tmp_path):
    store = _processed_store()
    assert [p.name for p in export_spectra_csv(store, tmp_path / "a.csv", data="raw")] == ["a_raw.csv"]
    assert [p.name for p in export_mzml(store, tmp_path / "a.mzml", data="processed")] == ["a_processed.mzml"]
    assert not (tmp_path / "a_processed.csv").exists()


def test_xlsx_workbook_has_raw_processed_peaks_and_round_trips(tmp_path):
    from openpyxl import load_workbook

    from ims_control.io.xlsx_io import export_xlsx, import_xlsx

    store = _processed_store()
    store.iterations[0].peaks = [
        {"index": 5, "time_ms": 1.0, "height": 2.0, "fwhm_ms": 0.1, "resolving_power": 10.0,
         "snr": 30.0, "k0": 1.5, "ccs_a2": None}
    ]
    path = export_xlsx(store, tmp_path / "run.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Raw", "Processed", "Peaks", "Settings"]
    wb.close()

    loaded = import_xlsx(path)
    _assert_raw_and_processed_preserved(loaded, store, exact=False)
    assert loaded.iterations[0].peaks[0]["ccs_a2"] is None
    assert loaded.iterations[1].peaks == []


def test_xlsx_selection(tmp_path):
    from openpyxl import load_workbook

    from ims_control.io.xlsx_io import export_xlsx, import_xlsx

    store = _processed_store()
    export_xlsx(store, tmp_path / "r.xlsx", data="raw")
    assert load_workbook(tmp_path / "r.xlsx").sheetnames == ["Raw", "Peaks", "Settings"]
    loaded = import_xlsx(tmp_path / "r.xlsx")
    np.testing.assert_allclose(
        loaded.iterations[0].raw_intensity, store.iterations[0].raw_intensity, rtol=1e-14, atol=0
    )


def test_invalid_selection_rejected(tmp_path):
    with pytest.raises(ValueError):
        export_hdf5(_processed_store(), tmp_path / "x.h5", data="everything")
