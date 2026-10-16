"""Processing/peak-detection settings and instrument metadata are recorded by every export."""

from __future__ import annotations

import json

import h5py
import numpy as np
import pytest
from openpyxl import load_workbook

from ims_control.acquisition.experiment import ExperimentConfig
from ims_control.io.hdf5_io import export_hdf5, import_hdf5
from ims_control.io.mzml_io import export_mzml, import_mzml
from ims_control.io.xlsx_io import SETTINGS_SHEET, export_xlsx, import_xlsx
from ims_control.models.data_store import DataStore
from ims_control.processing.settings import ProcessingSettings

SETTINGS = ProcessingSettings(
    positive_mode=True,
    subtract_baseline=True,
    baseline_window_ms=3.5,
    normalize=True,
    normalize_scale="0-100%",
    peak_min_height=0.25,
    peak_min_prominence=None,
    noise_window_start_ms=40.0,
    noise_window_end_ms=50.0,
)


@pytest.fixture
def store() -> DataStore:
    config = ExperimentConfig(
        num_points=50, exp_length_ms=25.0, drift_length_cm=12.5, drift_voltage_v=2345.0,
        pressure_torr=700.0, temperature_k=310.0, gas_type="He", ion_mz=150.5, ion_charge=2,
    )
    s = DataStore(config)
    rng = np.random.default_rng(1)
    for i in range(3):
        raw = rng.normal(size=50)
        s.add_iteration(i, raw * 2, raw)
    s.processing = SETTINGS
    return s


def _check(loaded: DataStore) -> None:
    assert loaded.processing == SETTINGS
    c = loaded.config
    assert (c.drift_length_cm, c.drift_voltage_v, c.pressure_torr, c.temperature_k) == (12.5, 2345.0, 700.0, 310.0)
    assert (c.gas_type, c.ion_mz, c.ion_charge) == ("He", 150.5, 2)


def test_hdf5_records_and_restores(store, tmp_path):
    path = tmp_path / "r.h5"
    export_hdf5(store, path)
    _check(import_hdf5(path))
    with h5py.File(path, "r") as f:
        assert f["processing"].attrs["baseline_window_ms"] == 3.5
        assert f["processing"].attrs["peak_min_prominence"] == "auto"
        assert f["instrument_metadata"].attrs["gas_type"] == "He"
        assert json.loads(f.attrs["processing_json"])["operations"]


def test_xlsx_settings_sheet(store, tmp_path):
    path = tmp_path / "r.xlsx"
    export_xlsx(store, path)
    _check(import_xlsx(path))
    ws = load_workbook(path)[SETTINGS_SHEET]
    params = {row[1] for row in ws.iter_rows(min_row=2, values_only=True)}
    assert {"ion_mz", "gas_type", "baseline_window_ms", "peak_min_prominence", "operations"} <= params


def test_mzml_records_and_restores(store, tmp_path):
    written = export_mzml(store, tmp_path / "r.mzML")
    text = written[0].read_text()
    assert 'name="processing:baseline_window_ms"' in text
    assert 'name="metadata:ion_mz"' in text
    _check(import_mzml(written[0]))


def test_legacy_file_without_settings_gets_defaults(store, tmp_path):
    path = tmp_path / "legacy.h5"
    export_hdf5(store, path)
    with h5py.File(path, "a") as f:
        del f.attrs["processing_json"]
    assert import_hdf5(path).processing == ProcessingSettings()
