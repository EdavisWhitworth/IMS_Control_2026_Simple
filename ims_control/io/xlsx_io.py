"""Excel (.xlsx) export/import: one workbook with Raw, Processed, and Peaks sheets.

.xlsx is used rather than legacy .xls, which is limited to 256 columns (one column per
iteration here) and 65,536 rows.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook

from ims_control.acquisition.experiment import ExperimentConfig
from ims_control.io.variants import DEFAULT_EXPORT, validate_export_data
from ims_control.models.data_store import DataStore

RAW_SHEET = "Raw"
PROCESSED_SHEET = "Processed"
PEAKS_SHEET = "Peaks"
PEAK_FIELDS = ["iteration", "index", "time_ms", "height", "fwhm_ms", "resolving_power", "snr", "k0", "ccs_a2"]


def export_xlsx(store: DataStore, path: str | Path, data: str = DEFAULT_EXPORT) -> Path:
    """Write `<path>` as a workbook. `data` is "both" (default), "raw" or "processed"; the
    selected spectra go on their own labelled sheets (one row per time point, one column per
    iteration) and the peak table is always on the Peaks sheet."""
    variants = validate_export_data(data)
    time_ms = store.time_axis_ms
    wb = Workbook()
    wb.remove(wb.active)
    for variant, sheet_name in (("raw", RAW_SHEET), ("processed", PROCESSED_SHEET)):
        if variant not in variants:
            continue
        ws = wb.create_sheet(sheet_name)
        ws.append(["time_ms"] + [f"iteration_{rec.index}" for rec in store.iterations])
        arrays = [rec.raw_intensity if variant == "raw" else rec.intensity for rec in store.iterations]
        for row_idx, t in enumerate(time_ms):
            ws.append([float(t)] + [float(a[row_idx]) for a in arrays])

    ws = wb.create_sheet(PEAKS_SHEET)
    ws.append(PEAK_FIELDS)
    for rec in store.iterations:
        for peak in rec.peaks:
            ws.append([rec.index] + [peak.get(k) for k in PEAK_FIELDS[1:]])

    path = Path(path)
    wb.save(path)
    return path


def _read_sheet(ws) -> tuple[np.ndarray, list[str], np.ndarray]:
    rows = list(ws.iter_rows(values_only=True))
    header, data_rows = rows[0], rows[1:]
    time_ms = np.array([float(r[0]) for r in data_rows])
    columns = np.array([[float(v) for v in r[1:]] for r in data_rows])
    return time_ms, [str(h) for h in header[1:]], columns


def import_xlsx(path: str | Path) -> DataStore:
    """Rebuild a DataStore from a workbook written by `export_xlsx`. Raw data is restored as
    `raw_intensity` when a Raw sheet exists; a processed-only workbook sets raw equal to the
    processed data, and a raw-only workbook uses raw for both (the GUI re-processes it).
    Instrument/ion metadata is not stored, so default metadata is used (edit it in the GUI)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    has_raw, has_proc = RAW_SHEET in wb.sheetnames, PROCESSED_SHEET in wb.sheetnames
    if not (has_raw or has_proc):
        raise ValueError("Workbook has neither a Raw nor a Processed sheet")
    time_ms, labels, proc = _read_sheet(wb[PROCESSED_SHEET if has_proc else RAW_SHEET])
    raw = None
    if has_raw:
        _, raw_labels, raw = _read_sheet(wb[RAW_SHEET])
        if raw_labels != labels or raw.shape != proc.shape:
            raise ValueError("Raw and Processed sheets do not describe the same iterations")

    num_points = len(time_ms)
    # exp_length_ms from sample spacing x count, matching DataStore.time_axis_ms (see csv_io).
    dt_ms = float(time_ms[1] - time_ms[0]) if num_points > 1 else 1.0
    store = DataStore(ExperimentConfig(exp_length_ms=dt_ms * num_points, num_points=num_points))
    for col, label in enumerate(labels):
        store.add_iteration(
            int(label.rsplit("_", 1)[-1]), proc[:, col], None if raw is None else raw[:, col]
        )

    if PEAKS_SHEET in wb.sheetnames:
        by_index = {rec.index: rec for rec in store.iterations}
        for row in list(wb[PEAKS_SHEET].iter_rows(values_only=True))[1:]:
            if row[0] is None or int(row[0]) not in by_index:
                continue
            peak = dict(zip(PEAK_FIELDS[1:], row[1:]))
            peak["index"] = int(peak["index"])
            by_index[int(row[0])].peaks.append(peak)
    wb.close()
    return store
